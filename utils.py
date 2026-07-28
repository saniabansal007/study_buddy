import os, re, json
import pypdf
from pptx import Presentation

import streamlit as st
from dotenv import load_dotenv
from langchain_huggingface import HuggingFaceEndpoint, ChatHuggingFace
from langchain_core.messages import SystemMessage, HumanMessage
from langchain_core.prompts import ChatPromptTemplate

from langchain_huggingface import HuggingFaceEmbeddings
from langchain_chroma import Chroma

from datetime import datetime, timedelta
import logging
logger = logging.getLogger("study_buddy")
logging.basicConfig(level=logging.INFO)

import math
import networkx as nx

load_dotenv()

# Constants
BASE_DIR   = os.path.abspath(os.path.dirname(__file__))
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")
CHROMA_DIR = os.path.join(BASE_DIR, "chroma_db")
CHUNK_SZ= 400
OVERLAP= 80
MAX_BYTES = 20 * 1024 * 1024   # 20 MB
MAX_CHARS = 10_000
HF_MODEL = "Qwen/Qwen2.5-7B-Instruct"
HF_TOKEN = os.environ.get("HF_TOKEN")
os.makedirs(UPLOAD_DIR, exist_ok=True)


# Returns true if the file has a supported extension
def ok_file(name: str)->bool:
    return "." in name and name.rsplit(".", 1)[1].lower() in {"pdf","pptx","txt"}


# File reading
def read_file(path: str, ext: str)->str:
    if ext== "pdf":
        reader= pypdf.PdfReader(path) 
        #reads all the pages and reader.pages returns [page1, page2, ...]

        return "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()
    
    elif ext== "pptx":
        prs= Presentation(path)
        parts= []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
        return "\n\n".join(parts).strip()
    
    else:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read().strip()


# Chunking 
def chunk(text: str)->list[str]:
    chunks, start= [], 0
    while start<len(text):
        end= start+ CHUNK_SZ
        chunks.append(text[start:end])
        start+= CHUNK_SZ- OVERLAP

    #using c.strip() discards the empty strings 
    return [c for c in chunks if c.strip()]

@st.cache_resource
def load_llm():
    try:
        endpoint = HuggingFaceEndpoint(
            repo_id=HF_MODEL,
            huggingfacehub_api_token=HF_TOKEN,
            max_new_tokens=1500,
            temperature=0.3,
        )
        llm = ChatHuggingFace(llm=endpoint)
        return llm
    except Exception as e:
        return None


def run_prompt(prompt_template: ChatPromptTemplate, variables: dict, max_tokens: int = 1500) -> str:
    chat_llm = load_llm()
    if chat_llm is None:
        raise RuntimeError("HF_TOKEN not set or LLM failed to initialise.")
    messages = prompt_template.format_messages(**variables)
    lc_msgs  = []
    for m in messages:
        if m.type == "system":
            lc_msgs.append(SystemMessage(content=m.content))
        else:
            lc_msgs.append(HumanMessage(content=m.content))
    return chat_llm.invoke(lc_msgs).content


SUMMARIZE_PROMPT= ChatPromptTemplate.from_messages([
    ("human",
     "Summarize the material below in {instr}. Use short headings and bullet points.\n\n"
     "MATERIAL:\n{text}")
])

def parse_json(raw: str):
    #Strip markdown fences the LLM wraps around JSON, then parse it.
    #Falls back to regex extraction if direct parsing fails.

    cleaned= re.sub(r"^```(?:json)?|```$", "", raw.strip(), flags=re.MULTILINE).strip()
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        for pat in [r"\[[\s\S]*\]", r"\{[\s\S]*\}"]:
            m = re.search(pat, cleaned)
            if m:
                try:
                    return json.loads(m.group(0))
                except:
                    pass
        raise


QUIZ_PROMPT= ChatPromptTemplate.from_messages([
    ("human",
     "Create {n} multiple-choice questions at {difficulty} difficulty.\n\n"
     "Return ONLY valid JSON:\n"
     '[{{"question":"...","topic":"...","options":["A","B","C","D"],"correct_index":0,"explanation":"..."}}]\n\n'
     "MATERIAL:\n{text}")
])


def sm2(card: dict, q: int) -> dict:
    """
    SM-2 spaced repetition algorithm.
    q = quality of recall:
        0 = Again (complete blank)
        3 = Hard  (correct but difficult)
        4 = Good  (correct with some effort)
        5 = Easy  (perfect recall)
    """
    q = max(0, min(5, q))          # clamp to valid range
    if q >= 3:                     # correct answer
        rep = card["repetitions"]
        if   rep == 0: iv = 1      # first time correct: review tomorrow
        elif rep == 1: iv = 6      # second time: review in 6 days
        else:          iv = round(card["interval"] * card["ease_factor"])
        card["repetitions"] += 1
    else:                          # wrong answer: reset
        card["repetitions"] = 0
        iv = 1
    ef= card["ease_factor"] + (0.1 - (5 - q) * (0.08 + (5 - q) * 0.02))
    card["ease_factor"] = round(max(1.3, ef), 3)   # never drop below 1.3
    card["interval"]    = iv
    card["next_review"] = (datetime.now() + timedelta(days=iv)).isoformat()
    card["history"].append({
        "quality":  q,
        "correct":  q >= 3,
        "interval": iv,
        "ts":       datetime.now().isoformat(),
    })
    return card


FLASHCARD_PROMPT= ChatPromptTemplate.from_messages([
    ("human",
     "Create {n} flashcard Q&A pairs from the material.\n\n"
     "Return ONLY valid JSON:\n"
     '[{{"front":"question","back":"detailed answer","topic":"topic name"}}]\n\n'
     "MATERIAL:\n{text}")
])


# CHAT FEATURE

@st.cache_resource
def load_embeddings():
    try:
        emb = HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2")
        return emb, True
    except Exception as e:
        logger.warning("Embeddings init failed: %s", e)
        return None, False


# module-level dict: doc_id → Chroma store instance
_chroma_stores: dict = {}


def _ensure_indexed(doc_id: str, chunks: list[str]):
    """Build the Chroma index for a document if it doesn't exist yet."""
    embeddings, RAG = load_embeddings()
    if not RAG or not chunks or doc_id in _chroma_stores:
        return
    try:
        store = Chroma.from_texts(
            texts=chunks,
            embedding=embeddings,
            collection_name=doc_id,
            persist_directory=os.path.join(CHROMA_DIR, doc_id),
        )
        _chroma_stores[doc_id] = store
        logger.info("Indexed %d chunks for %s", len(chunks), doc_id)
    except Exception as e:
        logger.error("_ensure_indexed: %s", e)


def retrieve(doc_id: str, query: str, k: int = 4) -> list[str]:
    """Return top-k relevant chunks for a query."""
    embeddings, RAG = load_embeddings()
    if not RAG:
        return []
    chunks = st.session_state.get("documents", {}).get(doc_id, {}).get("chunks", [])
    _ensure_indexed(doc_id, chunks)
    try:
        store = _chroma_stores.get(doc_id)
        if not store:
            return []
        return [d.page_content for d in store.similarity_search(query, k=k)]
    except Exception as e:
        logger.error("retrieve: %s", e)
        return []

def llm_invoke(messages: list, max_tokens: int = 1500) -> str:
    """Invoke the LLM with a plain list of dicts {role, content}."""
    chat_llm = load_llm()
    if chat_llm is None:
        raise RuntimeError("HF_TOKEN not set or LLM failed to initialise.")
    lc_messages = []
    for m in messages:
        role, content = m.get("role"), m.get("content", "")
        if role == "system":
            lc_messages.append(SystemMessage(content=content))
        else:
            lc_messages.append(HumanMessage(content=content))
    return chat_llm.invoke(lc_messages).content


def rag_answer(doc_id: str, question: str, fallback_text: str) -> tuple[str, list[str]]:
    _, rag_ready = load_embeddings()
    chunks = st.session_state.get("documents", {}).get(doc_id, {}).get("chunks", [])
    if rag_ready:
        _ensure_indexed(doc_id, chunks)
    if rag_ready and doc_id in _chroma_stores:
        try:
            sources = _chroma_stores[doc_id].similarity_search(question, k=4)
            context = "\n\n".join(d.page_content for d in sources)
            answer  = llm_invoke([{"role": "user", "content":
                f"Answer using ONLY the context below. Say if not in context.\n\n"
                f"CONTEXT:\n{context}\n\nQUESTION: {question}"}], max_tokens=800)
            return answer, [d.page_content for d in sources[:2]]
        except Exception as e:
            logger.error("rag_answer: %s", e)

    # fallback
    ctx = fallback_text[:MAX_CHARS]
    answer = llm_invoke([{"role": "user", "content":
        f"Answer using ONLY the context below. Say if not in context.\n\n"
        f"CONTEXT:\n{ctx}\n\nQUESTION: {question}"}], max_tokens=800)
    return answer, [ctx[:500]]




SOCRATIC_SYSTEM = (
    "You are a Socratic tutor. Ask guiding questions only- never give direct answers. "
    "One question at a time. Use this context:\n\n{context}"
)

GRAPH_PROMPT = ChatPromptTemplate.from_messages([
    ("human",
     "Extract 8-12 key concepts and their prerequisite relationships.\n\n"
     "Return ONLY valid JSON:\n"
     '{{"concepts":[{{"name":"...","description":"one sentence"}}],'
     '"edges":[{{"from":"prereq","to":"dependent","relation":"required for"}}]}}\n\n'
     "MATERIAL:\n{text}")
])

def svg_bar_chart(topics, accuracies, colors, W=740, H=280):
    if not topics:
        return '<p style="color:#756f5e;text-align:center;padding:40px">Complete a quiz to see topic accuracy.</p>'
    ml, mr, mt, mb = 48, 20, 24, 88
    cw  = W - ml - mr
    ch  = H - mt - mb
    n   = len(topics)
    bw  = cw / n * 0.6
    gap = cw / n
    p = [f'<svg viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg" style="width:100%;height:auto">']
    p.append(f'<rect width="{W}" height="{H}" fill="#ffffff"/>')
    for pct in [0, 25, 50, 75, 100]:
        y = mt + ch - pct / 100 * ch
        p.append(f'<line x1="{ml}" y1="{y:.1f}" x2="{W-mr}" y2="{y:.1f}" stroke="#e7e1d3" stroke-width="0.5"/>')
        p.append(f'<text x="{ml-5}" y="{y:.1f}" text-anchor="end" dominant-baseline="middle" font-size="10" font-family="Inter,sans-serif" fill="#756f5e">{pct}%</text>')
    for i, (t, a, c) in enumerate(zip(topics, accuracies, colors)):
        bx = ml + i * gap + (gap - bw) / 2
        bh = a / 100 * ch
        by = mt + ch - bh
        p.append(f'<rect x="{bx:.1f}" y="{by:.1f}" width="{bw:.1f}" height="{bh:.1f}" fill="{c}" rx="3"/>')
        p.append(f'<text x="{bx+bw/2:.1f}" y="{by-5:.1f}" text-anchor="middle" font-size="10" font-family="Inter,sans-serif" fill="#323a52">{a}%</text>')
        tx, ty = bx + bw / 2, mt + ch + 10
        label  = t[:13] + ("…" if len(t) > 13 else "")
        p.append(f'<text x="{tx:.1f}" y="{ty:.1f}" text-anchor="end" font-size="9" font-family="Inter,sans-serif" fill="#756f5e" transform="rotate(-35,{tx:.1f},{ty:.1f})">{label}</text>')
    p.append(f'<line x1="{ml}" y1="{mt}" x2="{ml}" y2="{mt+ch}" stroke="#e7e1d3" stroke-width="1"/>')
    p.append(f'<line x1="{ml}" y1="{mt+ch}" x2="{W-mr}" y2="{mt+ch}" stroke="#e7e1d3" stroke-width="1"/>')
    p.append('</svg>')
    return '\n'.join(p)


def svg_line_chart(fc_series, W=740, H=240):
    if not fc_series:
        return '<p style="color:#756f5e;text-align:center;padding:40px">Review flashcards to see interval growth.</p>'
    ml, mr, mt, mb = 45, 110, 20, 36
    cw     = W - ml - mr
    ch     = H - mt - mb
    all_iv = [iv for s in fc_series for iv in s["intervals"]]
    max_iv = max(all_iv, default=1) or 1
    max_x  = max((len(s["intervals"]) for s in fc_series), default=1)
    COLS   = ["#d9a441","#3f7d6e","#4a7fa5","#c1572a","#6a5acd","#2e8b57","#dc143c","#ff8c00"]
    p = [f'<svg viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg" style="width:100%;height:auto">']
    p.append(f'<rect width="{W}" height="{H}" fill="#ffffff"/>')
    for pct in [0, 25, 50, 75, 100]:
        val = pct / 100 * max_iv
        y   = mt + ch - (val / max_iv * ch)
        p.append(f'<line x1="{ml}" y1="{y:.1f}" x2="{W-mr}" y2="{y:.1f}" stroke="#e7e1d3" stroke-width="0.5"/>')
        p.append(f'<text x="{ml-4}" y="{y:.1f}" text-anchor="end" dominant-baseline="middle" font-size="9" font-family="Inter,sans-serif" fill="#756f5e">{round(val)}d</text>')
    for i, s in enumerate(fc_series[:8]):
        col = COLS[i % len(COLS)]
        ivs = s["intervals"]
        if not ivs:
            continue
        pts = [(ml + (j / max(max_x - 1, 1)) * cw, mt + ch - (iv / max_iv * ch)) for j, iv in enumerate(ivs)]
        if len(pts) > 1:
            d = "M " + " L ".join(f"{x:.1f},{y:.1f}" for x, y in pts)
            p.append(f'<path d="{d}" fill="none" stroke="{col}" stroke-width="2"/>')
        for x, y in pts:
            p.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="3" fill="{col}"/>')
        lx, ly = W - mr + 8, mt + i * 18
        label  = s["topic"][:11] + ("…" if len(s["topic"]) > 11 else "")
        p.append(f'<rect x="{lx}" y="{ly-5}" width="10" height="10" fill="{col}" rx="2"/>')
        p.append(f'<text x="{lx+14}" y="{ly:.1f}" dominant-baseline="middle" font-size="9" font-family="Inter,sans-serif" fill="#756f5e">{label}</text>')
    p.append(f'<line x1="{ml}" y1="{mt}" x2="{ml}" y2="{mt+ch}" stroke="#e7e1d3" stroke-width="1"/>')
    p.append(f'<line x1="{ml}" y1="{mt+ch}" x2="{W-mr}" y2="{mt+ch}" stroke="#e7e1d3" stroke-width="1"/>')
    p.append('</svg>')
    return '\n'.join(p)


def svg_concept_graph(nodes, edges_data, W=1000, H=600):
    if not nodes:
        return '<p style="color:#756f5e;text-align:center;padding:40px">No concepts extracted.</p>'
    G = nx.DiGraph()
    for n in nodes:
        G.add_node(n["id"] if "id" in n else n["name"])
    for n in nodes:
        if "id" not in n:
            n["id"] = n.get("name", "?")
    valid = [(e["from"], e["to"], e.get("relation", "")) for e in edges_data
             if e.get("from") in G.nodes and e.get("to") in G.nodes]
    G.add_edges_from([(u, v) for u, v, _ in valid])
    pos  = nx.spring_layout(G, seed=42, k=2.5 / math.sqrt(max(len(G.nodes), 1)))
    cent = nx.betweenness_centrality(G)
    mx   = max(cent.values(), default=0.001) or 0.001
    PAD  = 120

    def xy(p):
        return (round(PAD + (p[0] + 1) / 2 * (W - 2 * PAD), 1),
                round(PAD + (p[1] + 1) / 2 * (H - 2 * PAD), 1))

    coords = {n: xy(pos[n]) for n in G.nodes}
    p = [
        f'<svg viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg" '
        f'style="width:100%;height:{H}px;border:1px solid #e7e1d3;border-radius:10px;background:#fff;overflow:visible">',
        '<defs><marker id="arr" viewBox="0 -5 10 10" refX="20" refY="0" markerWidth="6" markerHeight="6" orient="auto">',
        '<path d="M0,-5L10,0L0,5" fill="#a8a08c"/></marker></defs>',
    ]
    for u, v, _ in valid:
        x1, y1 = coords[u]
        x2, y2 = coords[v]
        p.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="#c8c0b0" stroke-width="1.5" marker-end="url(#arr)"/>')
    for n in nodes:
        nid = n["id"]
        if nid not in coords:
            continue
        x, y = coords[nid]
        c    = cent.get(nid, 0)
        r    = round(18 + (c / mx) * 20, 1)
        fill = "#d9a441" if c > mx * 0.5 else "#1f2433"
        lbl  = nid
        desc = n.get("description", "").replace("<", "&lt;").replace(">", "&gt;")
        p.append(f'<g><title>{nid}: {desc}</title>')
        p.append(f'<circle cx="{x}" cy="{y}" r="{r}" fill="{fill}" stroke="white" stroke-width="2"/>')
        p.append(f'<text x="{x}" y="{y + r + 12}" text-anchor="middle" font-family="Inter,sans-serif" font-size="11" font-weight="600" fill="#1f2433">{lbl}</text>')
        p.append('</g>')
    p.append('</svg>')
    p.append('<p style="font-size:12px;color:#756f5e;margin-top:8px">Gold nodes = highest centrality · Arrow = prerequisite · Hover node for description</p>')
    return '\n'.join(p)